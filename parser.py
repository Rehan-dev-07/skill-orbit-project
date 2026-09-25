"""
Module 1: Resume Upload and Parsing
Handles document text extraction for PDF, DOCX, and TXT files.
Extracts contact metadata and detects resume section boundaries.
"""

import re
import os
from typing import Dict, Any, List
from pypdf import PdfReader
import docx

class ResumeParser:
    def __init__(self):
        # Email pattern
        self.email_pattern = re.compile(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+')
        # Phone pattern (supports international, indian numbers, dashes, dots, parens)
        self.phone_pattern = re.compile(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3,4}\)?[-.\s]?)?[\d\-.\s]{7,12}\d')
        # Links
        self.linkedin_pattern = re.compile(r'(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/([a-zA-Z0-9_-]+)', re.IGNORECASE)
        self.github_pattern = re.compile(r'(?:https?:\/\/)?(?:www\.)?github\.com\/([a-zA-Z0-9_-]+)', re.IGNORECASE)

        # Common resume section titles
        self.section_keywords = {
            "summary": ["summary", "professional summary", "career objective", "objective", "about me", "profile"],
            "education": ["education", "academic background", "qualifications", "academic credentials", "degrees"],
            "experience": ["experience", "work experience", "employment history", "professional experience", "internships", "work history"],
            "skills": ["skills", "technical skills", "core competencies", "skills & expertise", "tools & technologies", "technologies"],
            "projects": ["projects", "personal projects", "academic projects", "key projects", "work samples"],
            "certifications": ["certifications", "licenses", "certificates", "courses", "achievements", "awards"]
        }

    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extract plain text from PDF using pypdf."""
        text = ""
        try:
            reader = PdfReader(file_path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        except Exception as e:
            text = f"[PDF Extraction Error: {str(e)}]"
        return text.strip()

    def extract_text_from_docx(self, file_path: str) -> str:
        """Extract plain text from DOCX using python-docx."""
        text = []
        try:
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text.append(paragraph.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            text.append(cell.text)
        except Exception as e:
            text.append(f"[DOCX Extraction Error: {str(e)}]")
        return "\n".join(text).strip()

    def extract_text(self, file_path: str) -> str:
        """Detect file type and extract raw text."""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return self.extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return self.extract_text_from_docx(file_path)
        elif ext in [".txt", ".md"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read().strip()
        else:
            raise ValueError(f"Unsupported file format: {ext}. Only PDF, DOCX, and TXT are supported.")

    def extract_contact_info(self, text: str) -> Dict[str, Any]:
        """Extract contact details: email, phone, linkedin, github."""
        email_matches = self.email_pattern.findall(text)
        phone_matches = self.phone_pattern.findall(text)
        linkedin_matches = self.linkedin_pattern.findall(text)
        github_matches = self.github_pattern.findall(text)

        # Phone cleaning
        valid_phones = []
        for line in text.split("\n")[:25]:  # Look usually near top
            matches = list(re.finditer(r'(\+?\d{1,4}[-.\s]?)?(\(?\d{2,4}\)?[-.\s]?)?[\d\-.\s]{7,13}', line))
            for m in matches:
                p = m.group(0).strip()
                digits_count = sum(c.isdigit() for c in p)
                if 10 <= digits_count <= 14:
                    valid_phones.append(p)

        first_phone = valid_phones[0] if valid_phones else (phone_matches[0] if phone_matches else None)

        return {
            "email": email_matches[0] if email_matches else None,
            "phone": first_phone,
            "linkedin": f"linkedin.com/in/{linkedin_matches[0]}" if linkedin_matches else None,
            "github": f"github.com/{github_matches[0]}" if github_matches else None,
        }

    def detect_sections(self, text: str) -> Dict[str, bool]:
        """Detect presence of key resume sections."""
        lines = [line.strip().lower() for line in text.split("\n") if line.strip()]
        detected = {k: False for k in self.section_keywords.keys()}

        for line in lines:
            # Clean punctuation
            clean_line = re.sub(r'[^a-z\s]', '', line).strip()
            # If the line is short (under 40 chars), it's likely a section heading
            if len(clean_line) < 40:
                for section, triggers in self.section_keywords.items():
                    for trigger in triggers:
                        if trigger == clean_line or clean_line.startswith(trigger + " ") or clean_line.endswith(" " + trigger):
                            detected[section] = True

        return detected

    def parse(self, file_path: str, candidate_name: str = None) -> Dict[str, Any]:
        """Complete parsing pipeline."""
        raw_text = self.extract_text(file_path)
        word_count = len(raw_text.split())
        contact = self.extract_contact_info(raw_text)
        sections = self.detect_sections(raw_text)

        # Estimate candidate name from first non-empty lines if not supplied
        inferred_name = candidate_name
        if not inferred_name:
            lines = [l.strip() for l in raw_text.split("\n") if l.strip() and not l.startswith("[")]
            if lines:
                candidate = lines[0]
                # Filter out obvious headings or emails
                if len(candidate.split()) <= 4 and "@" not in candidate and not any(k in candidate.lower() for k in ["resume", "curriculum", "cv"]):
                    inferred_name = candidate
                else:
                    inferred_name = "Candidate"
            else:
                inferred_name = "Candidate"

        return {
            "file_name": os.path.basename(file_path),
            "candidate_name": inferred_name,
            "text": raw_text,
            "word_count": word_count,
            "contact": contact,
            "sections": sections
        }
